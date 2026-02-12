from __future__ import annotations

import os
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Tuple

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt


@dataclass
class BrandStyle:
    name: str
    primary: Tuple[int, int, int]
    secondary: Tuple[int, int, int]
    accent: Tuple[int, int, int]
    title_font: str
    body_font: str


STYLE_LIBRARY: Dict[str, BrandStyle] = {
    "luxury": BrandStyle(
        name="Noir Luxe",
        primary=(15, 15, 19),
        secondary=(42, 30, 24),
        accent=(215, 176, 106),
        title_font="Georgia",
        body_font="Aptos",
    ),
    "tech": BrandStyle(
        name="Neo Tech Premium",
        primary=(10, 20, 38),
        secondary=(25, 49, 88),
        accent=(125, 249, 255),
        title_font="Montserrat",
        body_font="Aptos",
    ),
    "fashion": BrandStyle(
        name="Editorial Vogue",
        primary=(30, 24, 26),
        secondary=(84, 58, 67),
        accent=(245, 216, 176),
        title_font="Bodoni MT",
        body_font="Aptos",
    ),
}


def infer_use_case(all_text: str) -> str:
    text = all_text.lower()
    if any(k in text for k in ["saas", "ai", "cloud", "platform", "api", "automation"]):
        return "tech"
    if any(k in text for k in ["fashion", "luxury", "collection", "style", "brand"]):
        return "fashion"
    return "luxury"


def pick_style(use_case: str, feedback: str) -> BrandStyle:
    feedback_lower = feedback.lower()
    if any(k in feedback_lower for k in ["tech", "futur", "cyber", "minimal"]):
        return STYLE_LIBRARY["tech"]
    if any(k in feedback_lower for k in ["fashion", "editorial", "runway", "magazine"]):
        return STYLE_LIBRARY["fashion"]
    return STYLE_LIBRARY.get(use_case, STYLE_LIBRARY["luxury"])


def _set_slide_background(slide, color: Tuple[int, int, int]):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*color)


def _style_text_frame(text_frame, style: BrandStyle, is_title: bool = False):
    for paragraph in text_frame.paragraphs:
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(10 if is_title else 4)
        for run in paragraph.runs:
            run.font.name = style.title_font if is_title else style.body_font
            run.font.bold = is_title
            run.font.size = Pt(40 if is_title else 20)
            run.font.color.rgb = RGBColor(*(style.accent if is_title else (236, 236, 236)))


def _text_density_score(words_per_slide: List[int]) -> float:
    if not words_per_slide:
        return 0.0
    overloaded = sum(1 for count in words_per_slide if count > 80)
    return round(max(0.0, 100 - (overloaded / len(words_per_slide) * 100)), 1)


def _visual_balance_score(image_count: int, slides: int) -> float:
    if slides == 0:
        return 0.0
    ratio = image_count / slides
    score = min(100.0, round(ratio * 100, 1))
    return max(score, 35.0)


def _estimate_brand_score(title_coverage: float, text_density: float, visual_balance: float) -> float:
    return round((title_coverage * 0.35) + (text_density * 0.3) + (visual_balance * 0.35), 1)


def _add_placeholder_image(slide, style: BrandStyle, temp_dir: Path) -> None:
    placeholder_path = temp_dir / f"placeholder-{uuid.uuid4().hex}.png"
    image = Image.new("RGB", (1280, 720), style.secondary)
    draw = ImageDraw.Draw(image)
    draw.rectangle((30, 30, 1250, 690), outline=style.accent, width=10)
    font = ImageFont.load_default()
    draw.text((420, 340), "Premium Placeholder Visual", fill=style.accent, font=font)
    image.save(placeholder_path)

    slide.shapes.add_picture(str(placeholder_path), Inches(6.0), Inches(1.25), Inches(7.0), Inches(3.9))


def _llm_ideas(summary: str) -> List[str]:
    api_key = os.getenv("GEMINI_API_KEY", "").strip()
    if not api_key:
        return []
    try:
        import google.generativeai as genai

        genai.configure(api_key=api_key)
        model = genai.GenerativeModel("gemini-1.5-flash")
        prompt = (
            "You are a commercial presentation design director. "
            "Return 3 concise recommendations to improve deck luxury styling and animation. "
            f"Deck summary: {summary}"
        )
        response = model.generate_content(prompt)
        lines = [line.strip("-• ") for line in response.text.splitlines() if line.strip()]
        return lines[:3]
    except Exception:
        return []


def analyze_presentation(input_path: str) -> Dict:
    prs = Presentation(input_path)
    all_text: List[str] = []
    image_count = 0
    title_count = 0
    words_per_slide: List[int] = []

    for slide in prs.slides:
        slide_words = 0
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            title_count += 1
            all_text.append(slide.shapes.title.text)
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                all_text.append(shape.text)
                slide_words += len(shape.text.split())
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_count += 1
        words_per_slide.append(slide_words)

    slides = len(prs.slides)
    combined = " ".join(all_text)
    use_case = infer_use_case(combined)
    title_coverage = round((title_count / max(slides, 1)) * 100, 1)
    text_density = _text_density_score(words_per_slide)
    visual_balance = _visual_balance_score(image_count, slides)
    brand_score = _estimate_brand_score(title_coverage, text_density, visual_balance)

    quality_flags = []
    if title_coverage < 70:
        quality_flags.append("Increase title consistency to create stronger narrative hierarchy.")
    else:
        quality_flags.append("Title hierarchy is consistent across slides.")
    if text_density < 70:
        quality_flags.append("Reduce text-heavy slides and convert content into visual blocks.")
    if image_count < max(1, slides // 2):
        quality_flags.append("Add more premium visuals to improve visual rhythm and brand storytelling.")

    summary = f"slides={slides}, use_case={use_case}, title_coverage={title_coverage}, brand_score={brand_score}"
    ai_ideas = _llm_ideas(summary)

    animation_plan = [
        "Hero title fade-up with slight scale (0.6s ease-out)",
        "Section transitions with cinematic cross-dissolve",
        "Staggered reveal for bullets/charts at 0.2s intervals",
    ]
    if ai_ideas:
        animation_plan.extend(ai_ideas)

    return {
        "slides": slides,
        "image_count": image_count,
        "title_coverage": title_coverage,
        "text_density": text_density,
        "visual_balance": visual_balance,
        "brand_score": brand_score,
        "use_case": use_case,
        "quality_flags": quality_flags,
        "animation_plan": animation_plan,
    }


def enhance_presentation(input_path: str, output_dir: str, feedback: str = "") -> Dict:
    report = analyze_presentation(input_path)
    style = pick_style(report["use_case"], feedback)

    prs = Presentation(input_path)
    temp_dir = Path(output_dir)
    temp_dir.mkdir(parents=True, exist_ok=True)
    modifications = []

    for index, slide in enumerate(prs.slides, start=1):
        _set_slide_background(slide, style.primary)
        modifications.append(f"Slide {index}: applied {style.name} background and premium palette")

        has_picture = False
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                has_picture = True
            if not getattr(shape, "has_text_frame", False):
                continue
            _style_text_frame(shape.text_frame, style, is_title=(shape == slide.shapes.title))

        if not has_picture:
            _add_placeholder_image(slide, style, temp_dir)
            modifications.append(f"Slide {index}: inserted premium placeholder visual")

        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = (
            "Motion cue: title fade-up (0.6s), key content stagger (0.2s delay), "
            "and subtle image parallax for cinematic delivery."
        )

    output_name = f"luxury-{uuid.uuid4().hex[:8]}.pptx"
    output_path = os.path.join(output_dir, output_name)
    prs.save(output_path)

    return {
        **report,
        "analysis": report,
        "style_applied": style.name,
        "fonts": [style.title_font, style.body_font],
        "modifications": modifications,
        "output_file": output_name,
    }
